"""
Análisis Big Five (OCEAN) de un perfil ya scrapeado, usando la API de Gemini (Google).

Lee lo que dejó scraper_instagram.py:
    datos/<usuario>/perfil.json
    datos/<usuario>/publicaciones.json
    datos/<usuario>/imagenes/...        (solo si usas --con-imagenes)
y guarda:
    datos/<usuario>/bigfive.json   (resultado estructurado + metadatos, para programas)
    datos/<usuario>/bigfive.txt    (informe legible: se abre en cualquier computador)
    datos/<usuario>/bigfive.xlsx   (Excel con tabla, gráfico y evidencias; requiere: pip install openpyxl)

Preparación (una sola vez):
    pip install google-genai
    La clave va en una variable de entorno, NUNCA escrita en el código ni subida a Git:
        Windows (PowerShell):  $env:GEMINI_API_KEY = "tu_clave"
        Linux / Mac:           export GEMINI_API_KEY="tu_clave"
    (en IntelliJ: Run > Edit Configurations > Environment variables)
    O, lo más cómodo: un archivo .env junto a este script (o en la carpeta desde donde lo ejecutas)
    con una sola línea:   GEMINI_API_KEY=tu_clave
    El script lo lee solo, por consola o con el botón Run. Agrega .env a tu .gitignore.

Uso:
    python analisis_bigfive_gemini.py mi_cuenta --dry-run          (muestra qué se enviaría; no llama a la API)
    python analisis_bigfive_gemini.py mi_cuenta                    (solo texto)
    python analisis_bigfive_gemini.py mi_cuenta --con-imagenes 8   (texto + hasta 8 imágenes)

Qué se envía a Google: usuario, biografía, conteos, texto de las publicaciones y, solo si lo
pides, algunas imágenes. NO se envían enlaces ni las listas de seguidores/seguidos.

Es una estimación exploratoria hecha con contenido público; no es una prueba psicométrica ni
un diagnóstico. Úsalo con tus propias cuentas o con personas que lo hayan consentido.
"""
import argparse
import datetime
import json
import os
import re
import sys
import textwrap

CARPETA_DATOS = "datos"
# Verifica el nombre vigente en https://ai.google.dev/gemini-api/docs/models y cámbialo con --modelo
MODELO_POR_DEFECTO = "gemini-3.5-flash"
MAX_PUBLICACIONES = 40
MAX_CARACTERES_POR_PUBLICACION = 600
MAX_IMAGENES = 10
MAX_BYTES_IMAGEN = 4_000_000
RASGOS = ["apertura", "responsabilidad", "extraversion", "amabilidad", "neuroticismo"]
NOMBRES = {"apertura": "Apertura a la experiencia", "responsabilidad": "Responsabilidad",
           "extraversion": "Extraversión", "amabilidad": "Amabilidad", "neuroticismo": "Neuroticismo"}
NIVELES = {"bajo", "medio", "alto"}
CONFIANZAS = {"baja", "media", "alta"}

INSTRUCCIONES = """Eres un asistente de investigación en psicología de la personalidad.
Recibirás datos públicos de un perfil de Instagram: biografía, conteos, el texto de sus
publicaciones y, a veces, algunas imágenes. Estima el perfil Big Five (OCEAN) usando
ÚNICAMENTE ese material.

Reglas:
- Rasgos (usa estas claves exactas, sin tildes): apertura, responsabilidad, extraversion,
  amabilidad, neuroticismo.
- Puntaje de 0 a 100 por rasgo, con nivel bajo (0-33), medio (34-66) o alto (67-100).
- Fundamenta cada rasgo con evidencia concreta, citando el id de la publicación (por ejemplo
  "post 3: ...") o la biografía. Las imágenes se identifican como "imagen del post N".
  No inventes evidencia.
- Si el material es escaso o ambiguo para un rasgo, baja la confianza y dilo.
- Es una estimación de cómo alguien se presenta en redes, que no equivale a cómo es.
  No hagas diagnósticos clínicos ni infieras salud mental, orientación sexual, religión,
  ideología política, origen étnico ni ningún otro dato sensible.
- Escribe en español. Responde SOLO con un objeto JSON válido, sin texto antes ni después ni
  bloques de código, con exactamente esta forma:
{
  "rasgos": {
    "apertura":        {"puntaje": 0, "nivel": "bajo|medio|alto", "confianza": "baja|media|alta", "evidencia": ["..."], "interpretacion": "..."},
    "responsabilidad": {...},
    "extraversion":    {...},
    "amabilidad":      {...},
    "neuroticismo":    {...}
  },
  "resumen": "párrafo breve",
  "limitaciones": ["..."]
}"""


# ---------------------------------------------------------------- configuración
def cargar_env():
    """Lee un archivo .env (carpeta actual o la del script) y define las variables que falten.

    Formato: CLAVE=valor, una por línea; se ignoran líneas vacías y comentarios (#).
    Una variable ya definida en el entorno con valor tiene prioridad; si está vacía, se rellena.
    """
    carpetas = [os.getcwd(), os.path.dirname(os.path.abspath(__file__))]
    for carpeta in dict.fromkeys(carpetas):
        ruta = os.path.join(carpeta, ".env")
        if not os.path.isfile(ruta):
            continue
        with open(ruta, encoding="utf-8-sig") as f:
            for linea in f:
                linea = linea.strip()
                if not linea or linea.startswith("#") or "=" not in linea:
                    continue
                clave, valor = linea.split("=", 1)
                clave = clave.replace("export ", "", 1).strip()
                valor = valor.strip().strip('"').strip("'")
                if clave and valor and not os.environ.get(clave):
                    os.environ[clave] = valor
        return ruta
    return None


# ---------------------------------------------------------------- entrada
def cargar_datos(usuario):
    base = os.path.join(CARPETA_DATOS, usuario)
    try:
        with open(os.path.join(base, "perfil.json"), encoding="utf-8") as f:
            perfil = json.load(f)
        with open(os.path.join(base, "publicaciones.json"), encoding="utf-8") as f:
            publicaciones = json.load(f)
    except FileNotFoundError as e:
        sys.exit(f"Falta {e.filename}. Primero ejecuta: python scraper_instagram.py {usuario}")
    return perfil, publicaciones


def recortar(texto, limite):
    texto = re.sub(r"\s+", " ", texto or "").strip()
    return texto if len(texto) <= limite else texto[:limite].rstrip() + "…"


def construir_material(perfil, publicaciones):
    """Texto que se envía a la API: solo lo necesario para el análisis."""
    partes = [
        f"Usuario: @{perfil.get('usuario')}",
        f"Biografía: {recortar(perfil.get('descripcion_perfil'), 500) or '(sin biografía)'}",
        f"Seguidores: {perfil.get('seguidores')} | Seguidos: {perfil.get('seguidos')} "
        f"| Publicaciones totales: {perfil.get('publicaciones_total')}",
        "",
        "Publicaciones (texto):",
    ]
    usadas = [p for p in publicaciones if (p.get("descripcion") or "").strip()][:MAX_PUBLICACIONES]
    for p in usadas:
        partes.append(f"[post {p.get('id')}] {recortar(p.get('descripcion'), MAX_CARACTERES_POR_PUBLICACION)}")
    if not usadas:
        partes.append("(no hay texto de publicaciones)")
    return "\n".join(partes), len(usadas)


def cargar_imagenes(usuario, publicaciones, cantidad):
    """Primera imagen de cada publicación, hasta 'cantidad'. Devuelve [(id_post, bytes)]."""
    base = os.path.join(CARPETA_DATOS, usuario)
    elegidas = []
    for p in publicaciones:
        if len(elegidas) >= min(cantidad, MAX_IMAGENES):
            break
        for img in p.get("imagenes") or []:
            ruta = os.path.join(base, img.get("archivo") or "")
            if os.path.isfile(ruta) and os.path.getsize(ruta) <= MAX_BYTES_IMAGEN:
                with open(ruta, "rb") as f:
                    elegidas.append((p.get("id"), f.read()))
                break
    return elegidas


# ---------------------------------------------------------------- API
def extraer_json(texto):
    """Toma el primer objeto JSON de la respuesta (tolera texto o ``` alrededor)."""
    inicio, fin = texto.find("{"), texto.rfind("}")
    if inicio == -1 or fin == -1:
        raise ValueError("La respuesta no contiene un objeto JSON.")
    return json.loads(texto[inicio:fin + 1])


def validar(resultado):
    """Comprueba que estén los cinco rasgos con puntaje, nivel y confianza válidos."""
    rasgos = resultado.get("rasgos", {})
    for r in RASGOS:
        d = rasgos.get(r)
        if not isinstance(d, dict):
            raise ValueError(f"Falta el rasgo '{r}'.")
        if not isinstance(d.get("puntaje"), (int, float)) or not 0 <= d["puntaje"] <= 100:
            raise ValueError(f"Puntaje inválido en '{r}'.")
        if d.get("nivel") not in NIVELES or d.get("confianza") not in CONFIANZAS:
            raise ValueError(f"Nivel o confianza inválidos en '{r}'.")
    return resultado


def llamar_modelo(material, imagenes, modelo):
    """Una llamada a la API de Gemini. Devuelve el texto de la respuesta."""
    try:
        from google import genai
        from google.genai import types
    except ImportError:
        sys.exit("Falta la librería: pip install google-genai")
    if not (os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")):
        sys.exit("No hay clave de API. Crea un archivo .env con la línea GEMINI_API_KEY=tu_clave "
                 "(junto al script o en la carpeta desde donde lo ejecutas) o define la variable de entorno.")

    contenido = [material]
    for id_post, datos in imagenes:
        contenido.append(f"Imagen del post {id_post}:")
        contenido.append(types.Part.from_bytes(data=datos, mime_type="image/jpeg"))

    try:
        cliente = genai.Client()   # lee GEMINI_API_KEY del entorno
        respuesta = cliente.models.generate_content(
            model=modelo,
            contents=contenido,
            config=types.GenerateContentConfig(
                system_instruction=INSTRUCCIONES,
                response_mime_type="application/json",   # pide la respuesta en JSON
            ),
        )
    except Exception as e:
        sys.exit(f"Error al llamar a Gemini: {e}\n"
                 "Revisa: clave válida, nombre del modelo (--modelo) y cuota (un 429 significa límite de uso: espera un poco).")
    return respuesta.text or ""


def analizar(material, imagenes, modelo, llamar=llamar_modelo, intentos=2):
    """Pide el análisis y lo valida; si la respuesta no cumple el formato, reintenta."""
    ultimo_error = None
    for n in range(1, intentos + 1):
        texto = llamar(material, imagenes, modelo)
        try:
            return validar(extraer_json(texto))
        except (ValueError, json.JSONDecodeError) as e:
            ultimo_error = e
            print(f"   !! Respuesta con formato inesperado (intento {n}/{intentos}): {e}")
    sys.exit(f"No se obtuvo un JSON válido: {ultimo_error}")


# ---------------------------------------------------------------- salida
def informe_texto(usuario, resultado, n_posts, n_imagenes, modelo):
    """Informe en texto plano, con líneas de hasta 90 caracteres."""
    def parrafo(texto, sangria="", continuacion=None):
        return textwrap.fill(texto or "", width=90, initial_indent=sangria,
                             subsequent_indent=sangria if continuacion is None else continuacion)

    meta = resultado.get("metadatos", {})
    lineas = [f"BIG FIVE ESTIMADO PARA @{usuario}", "=" * 90,
              f"Modelo: {modelo} | Publicaciones con texto: {n_posts} | Imágenes: {n_imagenes}"
              + (f" | Fecha: {meta['fecha']}" if meta.get("fecha") else ""),
              parrafo("Estimación exploratoria a partir de contenido público. No es un test psicométrico "
                      "ni un diagnóstico."),
              "", "RESULTADOS", "-" * 90,
              f"{'Rasgo':<28}{'Puntaje':>8}   {'Nivel':<8}{'Confianza':<10}"]
    for r in RASGOS:
        d = resultado["rasgos"][r]
        lineas.append(f"{NOMBRES[r]:<28}{d['puntaje']:>8}   {d['nivel']:<8}{d['confianza']:<10}")
    lineas += ["", "DETALLE POR RASGO", "-" * 90]
    for r in RASGOS:
        d = resultado["rasgos"][r]
        lineas += ["", f"{NOMBRES[r].upper()}  ({d['puntaje']}/100 - {d['nivel']}, confianza {d['confianza']})",
                   parrafo(d.get("interpretacion"), "  "), "  Evidencia:"]
        lineas += [parrafo(e, "    - ", "      ") for e in d.get("evidencia", [])]
    lineas += ["", "RESUMEN", "-" * 90, parrafo(resultado.get("resumen")), "", "LIMITACIONES", "-" * 90]
    lineas += [parrafo(l, "  - ", "    ") for l in resultado.get("limitaciones", [])]
    return "\n".join(lineas) + "\n"


def guardar_excel(ruta, usuario, resultado):
    """Excel con tres hojas: Resumen (tabla + gráfico), Evidencia y Metadatos."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("   (Para generar también el Excel: pip install openpyxl)")
        return False

    envolver = Alignment(wrap_text=True, vertical="top")
    cabecera_fuente = Font(bold=True, color="FFFFFF")
    cabecera_relleno = PatternFill("solid", fgColor="305496")

    def cabecera(hoja, fila, titulos):
        for c, t in enumerate(titulos, 1):
            celda = hoja.cell(fila, c, t)
            celda.font, celda.fill = cabecera_fuente, cabecera_relleno

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Big Five estimado para @{usuario}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Estimación exploratoria a partir de contenido público; no es un test psicométrico."
    fila0 = 4
    cabecera(ws, fila0, ["Rasgo", "Puntaje (0-100)", "Nivel", "Confianza", "Interpretación"])
    for i, r in enumerate(RASGOS, 1):
        d = resultado["rasgos"][r]
        ws.cell(fila0 + i, 1, NOMBRES[r]).alignment = Alignment(vertical="top")
        for col, valor in ((2, d["puntaje"]), (3, d["nivel"]), (4, d["confianza"])):
            ws.cell(fila0 + i, col, valor).alignment = Alignment(horizontal="center", vertical="top")
        interpretacion = ws.cell(fila0 + i, 5, d.get("interpretacion", ""))
        interpretacion.alignment = envolver
        ws.row_dimensions[fila0 + i].height = max(30, 15 * (len(d.get("interpretacion", "")) // 70 + 1))

    grafico = BarChart()
    grafico.type = "col"
    grafico.title = "Puntaje por rasgo (0-100)"
    grafico.add_data(Reference(ws, min_col=2, min_row=fila0, max_row=fila0 + len(RASGOS)), titles_from_data=True)
    grafico.set_categories(Reference(ws, min_col=1, min_row=fila0 + 1, max_row=fila0 + len(RASGOS)))
    grafico.y_axis.scaling.min, grafico.y_axis.scaling.max = 0, 100
    grafico.legend = None
    grafico.height, grafico.width = 8, 16
    ws.add_chart(grafico, f"A{fila0 + len(RASGOS) + 3}")

    fila = fila0 + len(RASGOS) + 20            # debajo del gráfico
    ws.cell(fila, 1, "Resumen").font = Font(bold=True)
    ws.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 1, end_column=5)
    ws.cell(fila + 1, 1, resultado.get("resumen", "")).alignment = envolver
    ws.row_dimensions[fila + 1].height = max(30, 15 * (len(resultado.get("resumen", "")) // 110 + 1))
    ws.cell(fila + 3, 1, "Limitaciones").font = Font(bold=True)
    for j, texto in enumerate(resultado.get("limitaciones", [])):
        ws.merge_cells(start_row=fila + 4 + j, start_column=1, end_row=fila + 4 + j, end_column=5)
        ws.cell(fila + 4 + j, 1, "- " + texto).alignment = envolver
        ws.row_dimensions[fila + 4 + j].height = max(18, 15 * (len(texto) // 110 + 1))
    for col, ancho in zip("ABCDE", (28, 16, 10, 12, 75)):
        ws.column_dimensions[col].width = ancho
    ws.page_setup.orientation = "landscape"          # al imprimir o exportar a PDF, todo cabe a lo ancho
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ev = wb.create_sheet("Evidencia")
    cabecera(ev, 1, ["Rasgo", "Evidencia citada"])
    fila = 2
    for r in RASGOS:
        for texto in resultado["rasgos"][r].get("evidencia", []):
            ev.cell(fila, 1, NOMBRES[r])
            ev.cell(fila, 2, texto).alignment = envolver
            fila += 1
    ev.column_dimensions["A"].width, ev.column_dimensions["B"].width = 28, 100

    meta = wb.create_sheet("Metadatos")
    cabecera(meta, 1, ["Dato", "Valor"])
    for k, v in resultado.get("metadatos", {}).items():
        meta.append([k, v])
    meta.column_dimensions["A"].width, meta.column_dimensions["B"].width = 28, 40

    wb.save(ruta)
    return True


def main():
    parser = argparse.ArgumentParser(description="Análisis Big Five de un perfil ya scrapeado (Gemini).")
    parser.add_argument("usuario", help="cuenta ya procesada con scraper_instagram.py")
    parser.add_argument("--modelo", default=MODELO_POR_DEFECTO, help=f"modelo (por defecto {MODELO_POR_DEFECTO})")
    parser.add_argument("--con-imagenes", type=int, default=0, metavar="N",
                        help=f"enviar también hasta N imágenes (máximo {MAX_IMAGENES}); por defecto solo texto")
    parser.add_argument("--dry-run", action="store_true", help="muestra lo que se enviaría y termina (sin API)")
    args = parser.parse_args()
    usuario = args.usuario.strip().lstrip("@").lower()
    cargar_env()

    perfil, publicaciones = cargar_datos(usuario)
    material, n_posts = construir_material(perfil, publicaciones)
    imagenes = cargar_imagenes(usuario, publicaciones, args.con_imagenes) if args.con_imagenes > 0 else []
    print(f"Publicaciones con texto: {n_posts} | Imágenes a enviar: {len(imagenes)} "
          f"| Caracteres de texto: {len(material)}")

    if args.dry_run:
        print("\n--- Esto es lo que se enviaría ---\n" + material)
        if imagenes:
            print("\n+ Imágenes de los posts: " + ", ".join(str(i) for i, _ in imagenes))
        return

    resultado = analizar(material, imagenes, args.modelo)
    resultado["metadatos"] = {
        "usuario": usuario, "modelo": args.modelo, "publicaciones_con_texto": n_posts,
        "imagenes_enviadas": len(imagenes),
        "fecha": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    base = os.path.join(CARPETA_DATOS, usuario)
    with open(os.path.join(base, "bigfive.json"), "w", encoding="utf-8") as f:
        json.dump(resultado, f, ensure_ascii=False, indent=4)
    with open(os.path.join(base, "bigfive.txt"), "w", encoding="utf-8") as f:
        f.write(informe_texto(usuario, resultado, n_posts, len(imagenes), args.modelo))
    hay_excel = guardar_excel(os.path.join(base, "bigfive.xlsx"), usuario, resultado)
    print(f"Listo: {base}/bigfive.json, {base}/bigfive.txt" + (f" y {base}/bigfive.xlsx" if hay_excel else ""))


if __name__ == "__main__":
    main()